"""Regressions from real papers (Scientific Reports PDF + SI DOCX), rebuilt synthetically."""

import io

import docx
import pytest

from manuscript_checker.analyze import analyze
from manuscript_checker.extract import Block, Document, load_document
from manuscript_checker.references import check_references
from manuscript_checker.structure import analyse_structure

canvas_mod = pytest.importorskip("reportlab.pdfgen.canvas")

BODY, SUP = ("Helvetica", 9), ("Helvetica", 6.3)


def _pdf(pages) -> bytes:
    """pages: list of lines; a line is a list of (text, kind) with kind in body|sup|bold."""
    buf = io.BytesIO()
    c = canvas_mod.Canvas(buf)
    for n, lines in enumerate(pages, 1):
        c.setFont(*BODY)
        c.drawString(40, 810, "www.example.com/journal/")
        y = 790
        for line in lines:
            x = 40
            for text, kind in line:
                font = {"body": BODY, "sup": SUP, "bold": ("Helvetica-Bold", 9)}[kind]
                c.setFont(*font)
                c.drawString(x, y + (3.5 if kind == "sup" else 0), text)
                x += c.stringWidth(text, *font)
            y -= 14
        c.setFont(*BODY)
        c.drawString(40, 30, f"Journal | (2025) 15:13504 | https://doi.org/10.1038/x {n}")
        c.showPage()
    c.save()
    return buf.getvalue()


def b(text):
    return [(text, "body")]


@pytest.fixture(scope="module")
def paper():
    pages = [
        [
            [("Alzheimer's disease (AD) is common among the aging population", "body"), ("1,2", "sup"), (".", "body")],
            [("Plaques form in the cortex", "body"), ("3–5", "sup"), (". We used ADMETlab 3.0", "body"), ("6", "sup"), (".", "body")],
            [("Simulations used GROMACS version 2018.1", "body"), ("7,8", "sup"), (" with a 10", "body"), ("5", "sup"),
             (" step limit and cells of 10 m", "body"), ("2", "sup"), (".", "body")],
            [("Fsp", "body"), ("3", "sup"), (" (fraction of sp", "body"), ("3", "sup"), (" carbons) and sp", "body"), ("3", "sup"),
             ("-hybridised atoms were counted; R", "body"), ("2", "sup"), (" was 0.9.", "body")],
            [("Toxicity was low for the phytochemicals (Fig.", "body")],
            [("S7). Binding was strong and stable when interacting with", "body")],
        ],
        [
            [("Fig. 2. Known inhibitors of BACE1 along with the selected phytochemicals", "body")],
            [("their optimal positioning explains the binding", "body"), ("9", "sup"), (".", "body")],
            [("Root mean square deviation (RMSD)", "bold")],
            [("The deviation stayed low for all complexes.", "body")],
            [("Figure S7. Toxicity profiles.", "body")],
            [("Figure 2 is discussed above. Volume 295(2) is an issue number, not a citation.", "body")],
            [("References", "bold")],
            [("1. Hampel, H. & Shen, Y. Beta-site amyloid precursor protein. J. Biol.", "body")],
            [("Chem. 5, 1-2 (2020). 2. Cole, S. L. & Vassar, R. Basic biology. Mol. Neurodegener.", "body")],
            [("2, 1-3 (2007). 3. Gohlke, H. Knowledge-based scoring. J. Mol. Biol. 295(2), 337-356 (2000).", "body")],
            [("4. Lee, K. Plaques. Cell 1, 2 (2019).", "body")],
            [("5. Kim, J. Cortex. Cell 2, 3 (2018).", "body")],
            [("6. Fu, L. et al. ADMETlab 3.0. Nucleic Acids Res. 52, 1 (2024).", "body")],
            [("7. Abraham, M. J. et al. GROMACS. SoftwareX 1, 19 (2015).", "body")],
            [("8. Hess, B. GROMACS 4. J. Chem. Theory Comput. 4, 435 (2008).", "body")],
            [("9. Liu, Y. Docking. Acta Pharmacol. Sin. 41, 138 (2020).", "body")],
        ],
        [b("Data availability"), b("All data are on Zenodo.")],
    ]
    return load_document("paper.pdf", _pdf(pages), "manuscript")


def test_running_headers_and_footers_are_removed(paper):
    assert not any("example.com" in bl.text or "15:13504" in bl.text for bl in paper.blocks)


def test_pdf_superscript_citations(paper):
    r = analyze([paper]).references
    assert r.style == "numeric" and r.detail == "superscript"
    assert [e.number for e in r.entries] == list(range(1, 10))
    # Every reference is cited exactly once: affiliation-free superscripts after words, a version
    # number (3.0) and a version-like year (2018.1) are citations; 10^5, m^2, sp^3, Fsp^3, R^2 are not.
    assert {e.number: len(e.citations) for e in r.entries} == {n: 1 for n in range(1, 10)}
    assert not r.unresolved and not r.issues


def test_reference_list_heading_and_wrapped_entries(paper):
    r = analyze([paper]).references
    assert r.entries[2].text.startswith("Gohlke") and "295(2)" in r.entries[2].text
    assert r.entries[0].text.endswith("(2020).")


def test_split_citation_and_page_top_caption(paper):
    a = analyze([paper])
    figures = {s.key.label: s.status for s in a.figures.items}
    assert figures["Figure 2"] == "OK"  # caption at the top of page 2 interrupts a sentence
    assert figures["Supplementary Figure 7 (S7)"] == "OK"  # "(Fig." / "S7)." across a line break


def test_bold_heading_is_its_own_block(paper):
    headings = [bl.text for bl in paper.blocks if bl.is_heading]
    assert "Root mean square deviation (RMSD)" in headings and "References" in headings


def test_bold_abstract_is_not_split_into_headings():
    lines = [[("This abstract is set in bold type across several lines of text,", "bold")],
             [("as some journals do, and none of these lines is a heading at all", "bold")],
             [("even though each line is short and entirely bold.", "bold")]]
    doc = load_document("a.pdf", _pdf([lines]), "manuscript")
    assert not any(bl.is_heading for bl in doc.blocks)


def test_glued_parenthesis_is_not_a_citation():
    blocks = [Block(t, "m.docx", "manuscript", i + 1) for i, t in enumerate(
        ["Shown (1) and (2) and (3) and (4) and (5), see J. Mol. Biol. 295(2) and Fsp(3).", "References",
         "1. A. B (2001).", "2. C. D (2002).", "3. E. F (2003).", "4. G. H (2004).", "5. I. J (2005)."])]
    r = check_references(analyse_structure(blocks))
    assert r.detail == "parenthesis"
    assert {e.number: len(e.citations) for e in r.entries} == {n: 1 for n in range(1, 6)}


def test_no_reference_list_is_one_error_not_one_per_number():
    blocks = [Block("Shown before [1, 2] and [3].", "m.docx", "manuscript", 1)]
    r = check_references(analyse_structure(blocks))
    assert [i.code for i in r.issues] == ["no_list"] and not r.unresolved


def test_docx_chemistry_superscripts_in_si_are_not_citations():
    d = docx.Document()
    p = d.add_paragraph("Descriptors: nRot, Fsp")
    p.add_run("3").font.superscript = True
    p.add_run(" (number of sp")
    p.add_run("3").font.superscript = True
    p.add_run(" hybridised carbons), R")
    p.add_run("2").font.superscript = True
    p.add_run(" = 0.9.")
    buf = io.BytesIO()
    d.save(buf)
    si = load_document("si.docx", buf.getvalue(), "supplementary")
    main = Document("m.docx", "manuscript", [Block(t, "m.docx", "manuscript", i + 1) for i, t in enumerate(
        ["Cited [1], [2] and [3].", "References", "1. A. B (2001).", "2. C. D (2002).", "3. E. F (2003)."])])
    r = analyze([main, si]).references
    assert all(c.block.source == "m.docx" for e in r.entries for c in e.citations)


def test_abstract_without_heading_ends_at_keywords():
    blocks = [Block(t, "m.pdf", "manuscript", i + 1) for i, t in enumerate([
        "Title of the paper",
        "Alzheimer's disease (AD) remains a challenge. " + "We studied inhibitors of the enzyme in detail. " * 12,
        "Keywords BACE1, docking",
        "Alzheimer's disease (AD) is a neurodegenerative disorder that affects millions. AD is common.",
    ])]
    a = analyze([Document("m.pdf", "manuscript", blocks)])
    assert a.proofing.stats["abstract_words"] and a.proofing.stats["abstract_words"] > 50
    notes = [f.note for c in a.proofing.checks if c.id == "abbreviations" for f in c.findings]
    assert not any(n.startswith("AD is defined again") for n in notes)
