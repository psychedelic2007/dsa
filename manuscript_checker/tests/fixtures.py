"""Realistic test manuscripts built with python-docx (no binary fixtures in the repo)."""

import io

import docx
from docx.oxml import parse_xml

W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'


def _sup(p, text):
    p.add_run(text).font.superscript = True


def _save(d) -> bytes:
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def numeric_manuscript() -> bytes:
    """Nature-style: superscript citations, author affiliations, Methods references."""
    d = docx.Document()
    d.add_paragraph("Allosteric control of binding kinetics", style="Title")
    p = d.add_paragraph("Jane Doe")
    _sup(p, "1,2")
    p.add_run(", John Roe")
    _sup(p, "2")
    p = d.add_paragraph("")
    _sup(p, "1")
    p.add_run("Department of Chemistry, Example University")
    d.add_paragraph("Abstract", style="Heading 1")
    d.add_paragraph("Surface plasmon resonance (SPR) is widely used to measure binding. We show that SPR kinetics "
                    "depend on allostery.")
    d.add_paragraph("Introduction", style="Heading 1")
    p = d.add_paragraph("Binding kinetics have been measured for decades")
    _sup(p, "1,2")
    p.add_run(". Isothermal titration calorimetry (ITC) complements SPR")
    _sup(p, "3–5")
    p.add_run(", and the cells were grown at 10 m")
    _sup(p, "2")
    p.add_run(" per flask (Fig. 1, Table 1). The the effect is large.")
    p = d.add_paragraph("We used ITC throughout (Fig. 2). Later work")
    _sup(p, "7")
    p.add_run(" extended this, as summarised in Table 2 and Error! Reference source not found.")
    p = d.add_paragraph("Before that, earlier reports")
    _sup(p, "6")
    p.add_run(" were inconclusive. Sample size was XX per group [ref].")
    ins = parse_xml(f'<w:ins {W} w:id="1" w:author="Reviewer" w:date="2026-01-01T00:00:00Z"><w:r><w:t> Added.</w:t></w:r></w:ins>')
    p._p.append(ins)
    d.add_paragraph("Table 1. Kinetic constants for all constructs.")
    t = d.add_table(rows=2, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Construct", "KD (nM)"
    t.rows[1].cells[0].text, t.rows[1].cells[1].text = "WT", "12"
    d.add_paragraph("Table 3. Uncited table.")
    d.add_paragraph("Methods", style="Heading 1")
    p = d.add_paragraph("Proteins were purified as described")
    _sup(p, "8")
    p.add_run(".")
    d.add_paragraph("References", style="Heading 1")
    refs = [
        "Smith, J. et al. Binding kinetics. Nature 570, 1–5 (2019). doi:10.1038/s41586-019-0001",
        "Lee, K. & Park, J. Calorimetry. Science 361, 10–12 (2018).",
        "Garcia, M. Allostery. Cell 170, 1–9 (2017).",
        "Chen, L. Kinetics II. J. Mol. Biol. 4, 1–2 (2016).",
        "Ito, H. Surface methods. Anal. Chem. 5, 3–4 (2015).",
        "Brown, A. Early reports. PNAS 100, 1–3 (2003).",
        "Novak, P. Later work. eLife 9, e1 (2020).",
        "Smith, J. et al. Binding kinetics. Nature 570, 1–5 (2019). doi:10.1038/s41586-019-0001",
    ]
    for i, r in enumerate(refs, 1):
        d.add_paragraph(f"{i}. {r}")
    d.add_paragraph("Acknowledgements", style="Heading 1")
    d.add_paragraph("We thank the facility staff.")
    d.add_paragraph("Data availability", style="Heading 1")
    d.add_paragraph("All data are available from the authors.")
    d.add_paragraph("Figure 1. Overview of binding.")
    d.add_paragraph("Figure 2. ITC traces.")
    data = d
    para = data.paragraphs[8]
    data.add_comment(para.runs[0], text="Check this", author="PI")
    return _save(data)


def author_year_manuscript() -> bytes:
    d = docx.Document()
    d.add_paragraph("Abstract", style="Heading 1")
    d.add_paragraph("Plant growth responds to light (Smith et al., 2019).")
    d.add_paragraph("Introduction", style="Heading 1")
    d.add_paragraph("Growth was shown before (Smith et al., 2019; Lee and Park 2018a, b). Müller (2020) disagreed, "
                    "and van der Waals (1873) explained it. Garcia et al. (2017) and (Jones, 2016) also (WHO, 2020).")
    d.add_paragraph("References", style="Heading 1")
    for r in [
        "Garcia, M. (2018). Wrong year. Science, 1, 1–2.",
        "Lee, K., & Park, J. (2018a). Title one. J Biol, 5, 1–10.",
        "Lee, K., & Park, J. (2018b). Title two. J Biol, 6, 1–10.",
        "Müller, H. (2020). Something. Cell, 1, 2.",
        "Smith, J. A., Jones, B., & Doe, C. (2019). Binding. Nature, 570, 1–5.",
        "van der Waals, J. D. (1873). Over de continuiteit. Leiden.",
        "World Health Organization. (2020). Report.",
        "Zed, A. (2010). Uncited work. Journal, 1, 1.",
    ]:
        d.add_paragraph(r, style="Bibliography") if "Bibliography" in [s.name for s in d.styles] else d.add_paragraph(r)
    d.add_paragraph("Funding", style="Heading 1")
    d.add_paragraph("Funded by a grant.")
    return _save(d)
