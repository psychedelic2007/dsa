import io

import docx
import pytest

from manuscript_checker.extract import Block, guess_role, load_blocks
from manuscript_checker.labels import TABLE, check_figures, check_labels, parse_caption, parse_mentions


def body(text, role="manuscript", style=""):
    return Block(text, "paper.docx", role, 1, style=style)


def labels(text, role="manuscript"):
    return [m.key.label for m in parse_mentions(body(text, role), in_caption=False)]


@pytest.mark.parametrize("text, expected", [
    ("as shown in Figure 2.", ["Figure 2"]),
    ("(Fig. 3a)", ["Figure 3"]),
    ("(Fig.3)", ["Figure 3"]),
    ("Figs. 1 and 4", ["Figure 1", "Figure 4"]),
    ("Figures 2, 3 and 5", ["Figure 2", "Figure 3", "Figure 5"]),
    ("Figures 2–4", ["Figure 2", "Figure 3", "Figure 4"]),
    ("Figures 2 to 4", ["Figure 2", "Figure 3", "Figure 4"]),
    ("Fig. 2a–d", ["Figure 2"]),  # panel range, not a figure range
    ("Fig. 1a–3b", ["Figure 1", "Figure 2", "Figure 3"]),
    ("Fig. 2a, b and 5c", ["Figure 2", "Figure 5"]),
    ("Figure 1(b)", ["Figure 1"]),
    ("Figures 1 and S2", ["Figure 1", "Supplementary Figure 2 (S2)"]),
    ("Figure S1–S3", ["Supplementary Figure 1 (S1)", "Supplementary Figure 2 (S2)", "Supplementary Figure 3 (S3)"]),
    ("Supplementary Figs. 2-3", ["Supplementary Figure 2 (S2)", "Supplementary Figure 3 (S3)"]),
    ("Suppl. Fig. 4", ["Supplementary Figure 4 (S4)"]),
    ("SI Figure 4", ["Supplementary Figure 4 (S4)"]),
    ("Extended Data Fig. 1", ["Extended Data Figure 1"]),
    ("Figure 3.2 shows", ["Figure 3.2"]),  # thesis chapter numbering
    ("Figures 3.1–3.3", ["Figure 3.1", "Figure 3.2", "Figure 3.3"]),
    ("see figure 2", ["Figure 2"]),
    ("Figure 2, 2019 data", ["Figure 2"]),  # years are not figure numbers
    ("Table 1 and Figure 6", ["Figure 6"]),
    ("Configured figures", []),
])
def test_mentions(text, expected):
    assert labels(text) == expected


def test_supplementary_file_unprefixed_numbers_are_supplementary():
    assert labels("Figure 1 shows", role="supplementary") == ["Supplementary Figure 1 (S1)"]


@pytest.mark.parametrize("text, style, expected", [
    ("Figure 1. Overview of the pipeline.", "", "Figure 1"),
    ("Fig. 2 | Cryo-EM map", "", "Figure 2"),
    ("**Figure 3:** Results", "", "Figure 3"),
    ("Figure 4 Distribution of scores", "", "Figure 4"),
    ("Figure 5", "", "Figure 5"),
    ("Supplementary Figure 2. Controls", "", "Supplementary Figure 2 (S2)"),
    ("Figure S7. Controls", "", "Supplementary Figure 7 (S7)"),
    ("Figure 3.2. Thesis chapter figure", "", "Figure 3.2"),
    ("Figure 6 shows the distribution", "Caption", "Figure 6"),  # Word caption style wins
    ("Figure 6 shows the distribution", "", None),  # body sentence, not a caption
])
def test_captions(text, style, expected):
    parsed = parse_caption(body(text, style=style))
    assert (parsed[0].key.label if parsed else None) == expected


def _report(*paragraphs, role="manuscript"):
    return check_figures([Block(t, "paper.docx", role, i + 1) for i, t in enumerate(paragraphs)])


def _status(report):
    return {f.key.label: f.status for f in report.items}


def test_end_to_end_statuses():
    report = _report(
        "Results are in Figure 1 and Fig. 3.",
        "Figure 1. First.",
        "Figure 2. Second, see also Fig. 1.",
        "Figure 4. Fourth, compare with Figure 2.",
    )
    assert _status(report) == {
        "Figure 1": "OK",
        "Figure 2": "UNCITED",  # only mentioned inside Figure 4's legend
        "Figure 3": "MISSING",
        "Figure 4": "UNCITED",
    }
    messages = " ".join(i.message for i in report.issues)
    assert "only mentioned inside other figure legends" in messages


def test_duplicate_caption_and_gap():
    report = _report("Figs. 1, 3.", "Figure 1. A.", "Figure 1. A again.", "Figure 3. C.")
    msgs = [i.message for i in report.issues]
    assert any("has 2 captions" in m for m in msgs)
    assert any("Numbering gap" in m and "Figure 2" in m for m in msgs)


def test_citation_order():
    report = _report("We start with Figure 2.", "Then Figure 1.", "Figure 1. A", "Figure 2. B")
    assert any("Figure 1 is first cited" in i.message for i in report.issues)


def test_list_of_figures_is_ignored():
    report = _report("Figure 1. Overview\t12", "Figure 1. Overview .......... 12", "See Figure 1.", "Figure 1. Overview")
    assert _status(report) == {"Figure 1": "OK"}
    assert not any("captions" in i.message for i in report.issues)


def test_legends_file_text_is_not_a_citation():
    blocks = [
        Block("Figure 1. A.", "legends.docx", "figures", 1),
        Block("(A) Same data as in Figure 2.", "legends.docx", "figures", 2),
        Block("Figure 2. B.", "legends.docx", "figures", 3),
        Block("We show Figure 1.", "paper.docx", "manuscript", 1),
    ]
    assert _status(check_figures(blocks)) == {"Figure 1": "OK", "Figure 2": "UNCITED"}


def test_supp_cited_only_in_supplement_is_flagged():
    blocks = [
        Block("Main text with no supplementary citations.", "paper.docx", "manuscript", 1),
        Block("Figure S1. Control.", "si.docx", "supplementary", 1),
        Block("As Figure S1 shows ...", "si.docx", "supplementary", 2),
    ]
    report = check_figures(blocks)
    assert _status(report) == {"Supplementary Figure 1 (S1)": "OK"}
    assert any(i.severity == "info" for i in report.issues)


def _docx_bytes(build):
    d = docx.Document()
    build(d)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def test_docx_reads_tables_and_caption_style():
    def build(d):
        d.add_paragraph("As seen in Fig. 1 and Figure 2.")
        cell = d.add_table(rows=1, cols=1).rows[0].cells[0]
        cell.paragraphs[0].text = "Figure 1. Caption inside a table cell."
        d.add_paragraph("Figure 2 shows nothing; it is a caption by style.", style="Caption")

    blocks = load_blocks("paper.docx", _docx_bytes(build), "manuscript")
    assert _status(check_figures(blocks)) == {"Figure 1": "OK", "Figure 2": "OK"}


def test_unsupported_types():
    with pytest.raises(ValueError, match="save it as .docx"):
        load_blocks("paper.doc", b"", "manuscript")
    with pytest.raises(ValueError, match="unsupported"):
        load_blocks("paper.tex", b"", "manuscript")


@pytest.mark.parametrize("name, role", [
    ("Smith_2026_manuscript.docx", "manuscript"),
    ("Figure_legends.docx", "figures"),
    ("Supplementary_Information.docx", "supplementary"),
    ("paper_SI.pdf", "supplementary"),
])
def test_guess_role(name, role):
    assert guess_role(name) == role


@pytest.mark.parametrize("numbered", [False, True])
def test_pdf_wrapped_citation_axis_labels_and_line_numbers(numbered):
    canvas = pytest.importorskip("reportlab.pdfgen.canvas")
    lines = [
        "The binding affinity was measured as shown in",
        "Figure 2. The values agree with theory.",  # wrapped citation, not a caption
        "Panels were compared (see",
        "Fig. 1a-c). Nothing else is cited here.",
        "Time (s)",  # axis label directly above a caption
        "Figure 1. Kinetics of binding.",
        "Figure 2 | Affinity measurements.",
        "Figure 3. Uncited figure.",
    ]
    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    for i, line in enumerate(lines, 1):
        c.drawString(40, 800 - 16 * i, (f"{i}   " if numbered else "") + line)
    c.save()
    report = check_figures(load_blocks("paper.pdf", buf.getvalue(), "manuscript"))
    assert _status(report) == {"Figure 1": "OK", "Figure 2": "OK", "Figure 3": "UNCITED"}
    assert all(len(f.captions) == 1 for f in report.items)


@pytest.mark.parametrize("caption", [
    "Figure S6. Title.",
    "Figure S6 (A) Distribution of scores.",
    "Figure S6 (a–c) Distribution of scores.",
    "Figure S6a. Title",
    "Figure S6. Title",
    "Figure S​6. Title",
    "Figure­S6. Title",
    "Figure S-6. Title",
    "Figure S6. Values 1 ... 10",
    "[Figure S6] Title",
])
def test_caption_variants_are_found(caption):
    from manuscript_checker.extract import clean
    blocks = [
        Block("We show Figure S6.", "m.docx", "manuscript", 1),
        Block(clean(caption), "si.docx", "supplementary", 1),
    ]
    assert _status(check_figures(blocks)) == {"Supplementary Figure 6 (S6)": "OK"}


def test_docx_line_break_before_caption():
    def build(d):
        d.add_paragraph("We show Figure S6.")
        p = d.add_paragraph("(a) (b)")
        p.add_run().add_break()
        p.add_run("Figure S6. Caption after a manual line break.")

    blocks = load_blocks("si.docx", _docx_bytes(build), "manuscript")
    assert _status(check_figures(blocks)) == {"Supplementary Figure 6 (S6)": "OK"}


def test_unrecognised_caption_is_pointed_out():
    blocks = [
        Block("We show Figure S6.", "m.docx", "manuscript", 1),
        Block("Figure S6 distribution of scores across runs", "si.docx", "supplementary", 4),
    ]
    report = check_figures(blocks)
    assert _status(report) == {"Supplementary Figure 6 (S6)": "MISSING"}
    hint = [i.message for i in report.issues if i.severity == "info"]
    assert hint and "si.docx (paragraph 4)" in hint[0]


def test_body_sentence_with_parenthetical_is_not_a_caption():
    assert parse_caption(body("Figure 2 (left) shows the trend.")) is None


def test_parenthetical_citation_is_not_a_suspected_caption():
    blocks = [Block("Mutational scanning (Fig. S5) supports the mechanism.", "m.docx", "manuscript", 1)]
    assert not any(i.code == "suspect" for i in check_figures(blocks).issues)
